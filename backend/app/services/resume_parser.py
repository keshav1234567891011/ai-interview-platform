import re
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from docx import Document
from fastapi import HTTPException
from pypdf import Configuration, PdfReader, apply_configuration

MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_TEXT = 100_000
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_PDF_STREAM = 8 * 1024 * 1024
PDF_LIMITS = Configuration(
    maximum_declared_stream_length=MAX_PDF_STREAM,
    array_based_stream_maximum_output_length=MAX_PDF_STREAM,
    zlib_maximum_output_length=MAX_PDF_STREAM,
    lzw_maximum_output_length=MAX_PDF_STREAM,
    run_length_maximum_output_length=MAX_PDF_STREAM,
    image_maximum_buffer_size=MAX_PDF_STREAM,
    jbig2_maximum_output_length=MAX_PDF_STREAM,
    jbig2dec_binary=None,  # Text parsing never searches for or launches an external image decoder.
    page_tree_maximum_entries=1000,
    page_tree_maximum_depth=30,
    xform_maximum_invocations_per_extraction=300,
)


def extract_pdf(data: bytes) -> str:
    # Context-local limits bound decompression before allocation; global settings stay unchanged.
    with apply_configuration(PDF_LIMITS):
        reader = PdfReader(BytesIO(data), strict=True)
        if reader.is_encrypted or len(reader.pages) > 50:
            raise ValueError("unsupported PDF")
        sections = []
        for page in reader.pages:
            stream = page.get_contents()
            if stream and len(stream.get_data()) > MAX_PDF_STREAM:
                raise ValueError("PDF content too large")
            sections.append(page.extract_text() or "")
            if sum(map(len, sections)) > MAX_TEXT:
                raise ValueError("too much text")
        return "\n".join(sections)


def validate_upload(filename: str, content_type: str | None, data: bytes) -> tuple[str, str]:
    name = re.split(r"[/\\]", filename)[-1]
    name = "".join(c for c in name if c.isprintable())[:180]
    extension = name.rsplit(".", 1)[-1].lower()
    expected = {"pdf": "application/pdf", "docx": DOCX_MIME}
    if extension not in expected or content_type not in {
        expected.get(extension),
        "application/octet-stream",
        None,
        "",
    }:
        raise HTTPException(415, "Upload a PDF or DOCX document.")
    if not data or len(data) > MAX_FILE_SIZE:
        raise HTTPException(413, "Upload a non-empty document no larger than 5 MB.")
    if (extension == "pdf" and not data.startswith(b"%PDF-")) or (
        extension == "docx" and not data.startswith(b"PK")
    ):
        raise HTTPException(422, "The document does not match its file format.")
    return name, expected[extension]


def extract_resume(data: bytes, content_type: str) -> str:
    try:
        if content_type == "application/pdf":
            text = extract_pdf(data)
        else:
            with ZipFile(BytesIO(data)) as archive:
                entries = archive.infolist()
                if (
                    len(entries) > 1000
                    or sum(item.file_size for item in entries) > 20 * 1024 * 1024
                    or "word/document.xml" not in archive.namelist()
                ):
                    raise BadZipFile("unsupported archive")
                if any(
                    "vbaproject" in item.filename.casefold() or item.flag_bits & 1
                    for item in entries
                ):
                    raise BadZipFile("active or encrypted content")
            document = Document(BytesIO(data))
            text = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [
                    cell.text
                    for table in document.tables
                    for row in table.rows
                    for cell in row.cells
                ]
            )
        text = text.strip()
        if len(text) > MAX_TEXT:
            raise ValueError("too much text")
        if not text:
            raise HTTPException(
                422, "No extractable text found. Scanned documents and OCR are not supported."
            )
        return text
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(
            422,
            "This document could not be read safely. "
            "Try an unencrypted PDF or DOCX with selectable text.",
        ) from None
